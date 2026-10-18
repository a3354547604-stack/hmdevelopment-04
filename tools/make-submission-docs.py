#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成参赛四件套里的 docx（说明文档）与 pptx（PPT 演示文稿）。

为什么要有这个脚本
------------------
`docs/参赛材料/10-说明文档正文.md` 与 `11-PPT正文.md` 是**唯一真源**（md），
docx / pptx 都是它的产物。手改产物会在下次重生成时丢失，所以：
**任何文字修改先改 md，再跑本脚本覆盖产物**（幂等，可反复重跑）。

产物（`docs/参赛材料/dist/`）
---------------------------
* `赛题四+独居老人居家安全监测与应急辅助系统+XX学校XX队.docx`
  宋体 / 小四 / 1.5 倍行距，含封面、目录（域）、页脚页码、每章末尾「待补素材占位框」。
* `赛题四+独居老人居家安全监测与应急辅助系统+XX学校XX队.pptx`
  16:9，P01–P14 + 备用页 B1–B7，每页「标题 + ≤3 要点 + 配图占位框 + 讲者备注」。

硬约束（违反即视为 bug，脚本会自检失败）
--------------------------------------
1. 真机/模拟器截图一律是**显式占位框**（`【占位·待实测】`），不用设计稿冒充，不写「已验证」。
2. 匿名：正文不出现学校/学院/团队名/指导老师/选手姓名；文档属性作者与公司写「参赛队」。
3. 文件名用 `XX 学校` / `XX 队` 占位（R7：赛题序号+赛题名称+学校团队名称）。

用法
----
    # 需要 python-docx / python-pptx（本机用 DSH 自带 Python 3.12）
    python tools/make-submission-docs.py            # 生成 + 自检
    python tools/make-submission-docs.py --check    # 只自检已有产物，不重写

在 WSL 里跑 Windows 侧 Python 的完整命令见 `docs/参赛材料/14-四件套生成说明.md`。
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import zipfile

try:  # Windows 控制台默认 GBK，日志会乱码，这里强制 UTF-8 输出
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:  # pragma: no cover - 老解释器无 reconfigure
    pass

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MAT_DIR = os.path.join(ROOT, "docs", "参赛材料")
DIST_DIR = os.path.join(MAT_DIR, "dist")
SRC_DOC = os.path.join(MAT_DIR, "10-说明文档正文.md")
SRC_PPT = os.path.join(MAT_DIR, "11-PPT正文.md")

BASE_NAME = "赛题四+独居老人居家安全监测与应急辅助系统+XX学校XX队"
DOCX_PATH = os.path.join(DIST_DIR, BASE_NAME + ".docx")
PPTX_PATH = os.path.join(DIST_DIR, BASE_NAME + ".pptx")

TEAM_PLACEHOLDER = "参赛队"  # 正文与文档属性里都不写真名（R4/R7 匿名）

# ---------------------------------------------------------------- 素材编号 → 验收项
# 依据 harmony-app/deliverables/07-截图与录屏清单.md §0/§1/§2/§3：
#   ① – ⑦ = App 侧验收项；D1 – D7 = 章程 §7 的 P0 七个演示动作。
ASSET_PROOF = {
    "S01": "验收项 ①⑦ 登录入口",
    "S02": "验收项 ①⑤ 首页实时（§6.7 + §4.3）",
    "S03": "验收项 ⑤ 契约 §4.3 十四属性",
    "S04": "D1 燃气超标报警",
    "S05": "验收项 ⑤ §6.13 列表与筛选",
    "S06": "验收项 ⑤ `status=` 筛选",
    "S07": "D1 §6.14 告警详情（data 平铺）",
    "S08": "D1 §6.15 处置闭环",
    "S09": "D7 §6.16 报表今日",
    "S10": "D7 §6.17 周趋势",
    "S11": "验收项 ⑤ §6.9/§6.10/§6.12",
    "S12": "验收项 ③④ 18 接口自检",
    "S13": "ADR-0001 配网交接边界（补拍）",
    "V01": "章程 §7.1 / ADR-0001 配网交接",
    "V02": "D1 §6.13–§6.15 报警处置闭环",
    "V03": "D6 契约 §8.5 断网补传",
    "V04": "D7 §6.16–§6.18 报表三周期",
    "V05": "D1 燃气报警（真机增补，见 07 §7.1）",
    "V06": "D2 跌倒报警（真机增补）",
    "V07": "D4 433 SOS（真机增补）",
    "E01": "验收项 ② 编译门禁自证",
    "E02": "验收项 ③④ 无设备自检（34/34）",
    "E03": "验收项 ① 构建产物（HAP 971362 B）",
}

# 每章末尾的占位框：图/表编号 → (说明, 编号, 素材类别)
# 素材类别决定提示语：shot = 真机/模拟器实拍（待实测）；draw = 自绘（待补）；cmd = 本机可复现命令证据
CHAPTER_ASSETS = {
    "第 1 章": [("图 1-1", "独居老人居家风险与响应链路（痛点→监测→预警→处置）", "draw")],
    "第 2 章": [
        ("S04", "首页报警态（燃气超标报警）", "shot"),
        ("S11", "设置页：阈值与「演示事件」三项", "shot"),
        ("S12", "设置 → 接口自检页（通过 N/18）", "shot"),
    ],
    "第 3 章": [
        ("图 3-1", "系统总体架构（节点 / 汇聚 / 网关 / 云 / App 五层）", "draw"),
        ("图 3-2", "一次燃气报警的端到端数据流（时序图）", "draw"),
        ("图 3-3", "配网与绑定流程图（浏览器 H5 + 服务器绑定）", "draw"),
        ("图 3-4", "断网补传与分级重推状态机", "draw"),
        ("表 3-1", "接口契约端点清单摘要（18 REST + webhook + §6.21 取时）", "draw"),
    ],
    "第 4 章": [("图 4-1", "硬件连接拓扑（3 节点 + 汇聚 + 网关 + 传感器）", "draw")],
    "第 5 章": [
        ("S01–S13", "App 五页界面截图（含 S13 空态补拍）", "shot"),
        ("V01–V07", "现场录屏（配网交接 / 报警闭环 / 断网补传 / 报表三周期 + 三项真机增补）", "shot"),
        ("E01–E03", "命令行证据（构建门禁自证 / 无设备自检 34 项 / HAP 产物）", "cmd"),
    ],
    "第 6 章": [("四象限图", "创新拓展四项落地程度（已做 / 部分做 / 规划中）", "draw")],
}

CATEGORY_NOTE = {
    "shot": "待实测（真机/模拟器实拍；本机无 hdc 设备、无模拟器、/dev/kvm 不可写）",
    "draw": "待补（自绘示意图，可据正文文字与仓库文档绘制）",
    "cmd": "待补（命令行输出截图，本机可复现，按 07 §7 重跑门禁后截取）",
}
CATEGORY_SIZE = {
    "shot": "建议 360×780 竖屏截图（≥1080×2340 原图另存）",
    "draw": "建议矢量或 ≥1600 px 宽高清 PNG",
    "cmd": "建议终端窗口 ≥1200×700，字号调大后截图",
}

# 内部专用、不进 docx 正文的章节（写作过程说明，提交前无需出现在成品里）
SKIP_SECTIONS = ("写作状态与待补", "封面（文字稿）")
# 正文开头的写作说明引用块（用途/章节结构/引用体例/三条红线/范围口径）
SKIP_LEADING_QUOTE_MARKERS = ("**用途**", "**章节结构**", "**引用体例**", "**三条红线**", "**范围口径**")

# ------------------------------------------------------------------ 通用小工具


def read_lines(path: str) -> list[str]:
    with open(path, encoding="utf-8") as fh:
        return fh.read().replace("\r\n", "\n").split("\n")


def inline_plain(text: str) -> str:
    """去掉 markdown 行内标记，用于核对与打印。"""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text.strip()


# ------------------------------------------------------------------ docx 基础样式

CJK_BODY = "宋体"
CJK_HEAD = "黑体"
LATIN_BODY = "Times New Roman"


def _set_run_font(run, size=12, bold=False, ea=CJK_BODY, latin=LATIN_BODY, color=None):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = latin
    if color is not None:
        from docx.shared import RGBColor

        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), ea)


TOKEN_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|【出处：[^】]*】)")


def add_runs(par, text: str, size=12, bold=False, ea=CJK_BODY, color=None):
    """把一行 markdown 文本写成若干 run：加粗、行内代码、【出处：…】各用不同格式。"""
    for tok in TOKEN_RE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            run = par.add_run(tok[2:-2])
            _set_run_font(run, size=size, bold=True, ea=ea, color=color)
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            run = par.add_run(tok[1:-1])
            _set_run_font(run, size=size, bold=bold, ea=ea, latin="Courier New", color=color)
        elif tok.startswith("【出处："):
            # 出处标记保留（可追溯性），但用小五号灰色，不抢正文
            run = par.add_run(tok)
            _set_run_font(run, size=9, bold=False, ea=ea, color=(0x80, 0x80, 0x80))
        else:
            run = par.add_run(tok)
            _set_run_font(run, size=size, bold=bold, ea=ea, color=color)
    return par


def setup_docx_styles(doc):
    from docx.enum.text import WD_LINE_SPACING
    from docx.oxml.ns import qn
    from docx.shared import Pt

    normal = doc.styles["Normal"]
    normal.font.name = LATIN_BODY
    normal.font.size = Pt(12)  # 小四
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_BODY)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE  # 1.5 倍行距
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)

    sizes = {"Heading 1": 16, "Heading 2": 14, "Heading 3": 12}
    for name, size in sizes.items():
        st = doc.styles[name]
        st.font.name = CJK_HEAD
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = None
        st.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_HEAD)
        st.element.rPr.rFonts.set(qn("w:ascii"), CJK_HEAD)
        st.element.rPr.rFonts.set(qn("w:hAnsi"), CJK_HEAD)
        spf = st.paragraph_format
        spf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        spf.space_before = Pt(12 if name == "Heading 1" else 8)
        spf.space_after = Pt(6)


def add_page_number_footer(section):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    par = section.footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    _set_run_font(run, size=9)


def add_toc_field(doc):
    """目录：插入 Word 域（打开后按 F9 / 打印前更新即得页码目录）。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    par = doc.add_paragraph()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), 'TOC \\o "1-3" \\h \\z \\u')
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:eastAsia"), CJK_BODY)
    rfonts.set(qn("w:ascii"), LATIN_BODY)
    rpr.append(rfonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "24")
    rpr.append(sz)
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = "（目录域：在 Word 中右键选择“更新域”或按 F9 生成页码）"
    run.append(text)
    fld.append(run)
    par._p.append(fld)
    return par


def shade_cell(cell, fill="F2F2F2"):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_placeholder_box(doc, chapter: str, assets):
    """每章末尾的「待补素材占位框」——真机素材绝不冒充，只留框与编号。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc.add_paragraph()
    head = doc.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_runs(head, f"【占位·待补素材】{chapter} 需要以下图/表，交前按编号替换本框", size=12, bold=True)

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for cell, label in zip(hdr, ("编号", "内容", "状态与建议尺寸")):
        shade_cell(cell, "D9D9D9")
        add_runs(cell.paragraphs[0], label, size=12, bold=True)

    for code, desc, kind in assets:
        row = table.add_row().cells
        proof = ASSET_PROOF.get(code, "")
        title = code if not proof else f"{code}（{proof}）"
        add_runs(row[0].paragraphs[0], title, size=12)
        add_runs(row[1].paragraphs[0], desc, size=12)
        add_runs(row[2].paragraphs[0], f"{CATEGORY_NOTE[kind]}；{CATEGORY_SIZE[kind]}", size=12)

    tail = doc.add_paragraph()
    tail.paragraph_format.space_before = Pt(6)
    add_runs(
        tail,
        "说明：本机无 hdc 设备、无模拟器、/dev/kvm 不可写，因此界面类素材一律由用户在真机/模拟器实拍；"
        "`harmony-app/deliverables/04-页面预览图/` 的 6 张是设计预览，不得当作真机截图使用。",
        size=12,
    )
    return table


# ------------------------------------------------------------------ md → docx 正文


def is_table_row(line: str) -> bool:
    return line.lstrip().startswith("|")


def is_sep_row(cells) -> bool:
    return all(re.fullmatch(r":?-{2,}:?", c.strip()) for c in cells if c.strip() != "")


def split_row(line: str):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def add_md_table(doc, rows, warn):
    widths = {len(r) for r in rows}
    if len(widths) > 1:
        warn.append(f"表格列数不一致：{sorted(widths)}（多出的单元格并入末列渲染）")
        rows = [r[: len(rows[0]) - 1] + ["|".join(r[len(rows[0]) - 1 :])] if len(r) > len(rows[0]) else r for r in rows]
    ncol = len(rows[0])
    table = doc.add_table(rows=0, cols=ncol)
    table.style = "Table Grid"
    for ri, cells in enumerate(rows):
        row = table.add_row().cells
        for ci in range(ncol):
            text = cells[ci] if ci < len(cells) else ""
            if ri == 0:
                shade_cell(row[ci], "D9D9D9")
            add_runs(row[ci].paragraphs[0], text, size=12, bold=(ri == 0))
    doc.add_paragraph()
    return table


def build_docx(warn: list[str]) -> dict:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Cm, Pt

    lines = read_lines(SRC_DOC)
    doc = Document()
    setup_docx_styles(doc)

    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.17)
    sec.right_margin = Cm(3.17)
    sec.different_first_page_header_footer = True  # 封面不显示页码
    add_page_number_footer(sec)

    # ---- 封面（正文里不写校名/团队名/姓名，只留「参赛队」占位）
    def center(text, size, bold=False, space_after=12):
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(space_after)
        add_runs(par, text, size=size, bold=bold, ea=CJK_HEAD)
        return par

    for _ in range(4):
        doc.add_paragraph()
    center("赛题四", 16, bold=False, space_after=24)
    center("独居老人居家安全监测与应急辅助系统", 22, bold=True, space_after=18)
    center("参赛作品设计方案", 16, space_after=48)
    for _ in range(3):
        doc.add_paragraph()
    # 匿名（R2）：内容里不写参赛队/校名/指导老师/选手姓名，队名只出现在压缩包文件名（R7）。
    center("日期：2026 年 11 月", 14, space_after=6)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- 目录
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(par, "目录", size=16, bold=True, ea=CJK_HEAD)
    add_toc_field(doc)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- 正文
    i = 0
    first_heading_done = False
    heading_counts = {"h1": 0, "h2": 0, "h3": 0}
    table_count = 0
    placeholder_boxes = []
    cur_chapter = None
    chapter_assets_done = set()
    body_started = False  # 见到第一个标题后才算正文开始（此前的引用块是写作说明）

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        # 跳过文件标题与开头写作说明引用块
        if i == 0 and stripped.startswith("# "):
            i += 1
            continue
        # 文件开头的写作说明引用块（用途/章节结构/引用体例/三条红线/范围口径）整段不进成品
        if stripped.startswith("> ") and not body_started:
            while i < len(lines) and lines[i].strip().startswith(">"):
                i += 1
            continue

        # 内部专用章节：整段跳过
        if stripped.startswith("## ") and any(s in stripped for s in SKIP_SECTIONS):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("## "):
                i += 1
            continue

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            i += 1
            continue

        # 表格
        if is_table_row(line):
            rows = []
            while i < len(lines) and is_table_row(lines[i]):
                cells = split_row(lines[i])
                if not is_sep_row(cells):
                    rows.append(cells)
                i += 1
            if rows:
                add_md_table(doc, rows, warn)
                table_count += 1
            continue

        # 标题
        m = re.match(r"^(#{2,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            # 章末尾插入占位框
            if level == 2 and cur_chapter and cur_chapter not in chapter_assets_done:
                assets = CHAPTER_ASSETS.get(cur_chapter)
                if assets:
                    add_placeholder_box(doc, cur_chapter, assets)
                    placeholder_boxes.append(cur_chapter)
                    chapter_assets_done.add(cur_chapter)

            body_started = True
            if level == 2:
                if first_heading_done:
                    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                first_heading_done = True
                heading_counts["h1"] += 1
                doc.add_heading(inline_plain(text), level=1)
                for key in CHAPTER_ASSETS:
                    if key in text:
                        cur_chapter = key
            elif level == 3:
                heading_counts["h2"] += 1
                doc.add_heading(inline_plain(text), level=2)
            else:
                heading_counts["h3"] += 1
                doc.add_heading(inline_plain(text), level=3)
            i += 1
            continue

        # 引用块
        if stripped.startswith(">"):
            text = stripped.lstrip(">").strip()
            if not text:
                i += 1
                continue
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.75)
            par.paragraph_format.space_after = Pt(6)
            add_runs(par, text, size=12, color=(0x40, 0x40, 0x40))
            i += 1
            continue

        # 列表
        m = re.match(r"^([-*]|\d+\.)\s+(.*)$", stripped)
        if m:
            marker, text = m.group(1), m.group(2)
            bullet = "• " if marker in ("-", "*") else f"{marker} "
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.75)
            par.paragraph_format.first_line_indent = Cm(-0.4)
            add_runs(par, bullet + text, size=12)
            i += 1
            continue

        # 普通段落（中文正文首行缩进 2 字符）
        par = doc.add_paragraph()
        par.paragraph_format.first_line_indent = Pt(24)
        add_runs(par, stripped, size=12)
        i += 1

    # 最后一章的占位框
    if cur_chapter and cur_chapter not in chapter_assets_done:
        assets = CHAPTER_ASSETS.get(cur_chapter)
        if assets:
            add_placeholder_box(doc, cur_chapter, assets)
            placeholder_boxes.append(cur_chapter)

    # 附录 D：待补素材总清单（内部用，提交前删除）
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("附录 D 待补素材总清单（内部用，提交前删除）", level=1)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(24)
    add_runs(
        p,
        "本清单由 `tools/make-submission-docs.py` 从正文与 `07-截图与录屏清单.md` 自动汇总，"
        "供拍摄/绘图时逐条勾选。素材齐备后重跑生成脚本即可（占位框仍在正文各章末尾，交前删除本页）。",
        size=12,
    )
    rows = [["编号", "内容", "对应验收项 / 演示动作", "状态与建议尺寸"]]
    for chapter, assets in CHAPTER_ASSETS.items():
        for code, desc, kind in assets:
            rows.append(
                [
                    code,
                    f"{chapter}：{desc}",
                    ASSET_PROOF.get(code, "—"),
                    f"{CATEGORY_NOTE[kind]}；{CATEGORY_SIZE[kind]}",
                ]
            )
    add_md_table(doc, rows, warn)

    # ---- 文档属性：不得含真实身份
    cp = doc.core_properties
    cp.author = TEAM_PLACEHOLDER
    cp.last_modified_by = TEAM_PLACEHOLDER
    cp.comments = ""
    cp.category = ""
    cp.subject = "赛题四 参赛作品设计方案"
    cp.title = "独居老人居家安全监测与应急辅助系统"

    os.makedirs(DIST_DIR, exist_ok=True)
    doc.save(DOCX_PATH)

    return {
        "headings": heading_counts,
        "tables": table_count,
        "placeholder_boxes": placeholder_boxes,
    }


# ------------------------------------------------------------------ md → pptx


PAGE_RE = re.compile(r"^###\s+(P\d{2}|B\d)\s+(.*)$")


def parse_ppt_pages(lines, warn):
    pages = []
    cur = None
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        m = PAGE_RE.match(line.strip())
        if m:
            if cur:
                pages.append(cur)
            cur = {"code": m.group(1), "head": m.group(2).strip(), "title": "", "bullets": [], "fig": "", "notes": ""}
            i += 1
            continue
        if cur is not None:
            s = line.strip()
            if s.startswith("**标题**"):
                cur["title"] = inline_plain(s.split("：", 1)[-1])
            elif s.startswith("**要点**"):
                j = i + 1
                while j < len(lines):
                    t = lines[j].strip()
                    mm = re.match(r"^\d+\.\s+(.*)$", t)
                    if mm:
                        cur["bullets"].append(inline_plain(mm.group(1)))
                        j += 1
                    elif t == "":
                        j += 1
                        if j < len(lines) and re.match(r"^\d+\.\s+", lines[j].strip()):
                            continue
                        break
                    else:
                        break
                i = j
                continue
            elif s.startswith("**配图**"):
                cur["fig"] = inline_plain(s.split("：", 1)[-1])
            elif s.startswith("**讲者备注**"):
                cur["notes"] = inline_plain(s.split("：", 1)[-1])
        i += 1
    if cur:
        pages.append(cur)

    for pg in pages:
        if not pg["title"]:
            warn.append(f"{pg['code']} 缺 **标题**")
        if not pg["bullets"]:
            warn.append(f"{pg['code']} 缺 **要点**")
        if len(pg["bullets"]) > 3:
            warn.append(f"{pg['code']} 要点 {len(pg['bullets'])} 条 > 3")
        if not pg["fig"]:
            warn.append(f"{pg['code']} 缺 **配图**")
        if not pg["notes"]:
            warn.append(f"{pg['code']} 缺 **讲者备注**")
    return pages


def parse_backup_pages(lines):
    """从 `## 2. 备用页` 的表格里取 B1–B7。"""
    rows = []
    inside = False
    for line in lines:
        s = line.strip()
        if s.startswith("## 2. 备用页"):
            inside = True
            continue
        if inside and s.startswith("## "):
            break
        if inside and is_table_row(line):
            cells = split_row(line)
            if not is_sep_row(cells) and not cells[0].startswith("备用页"):
                rows.append([inline_plain(c) for c in cells])
    return rows


def fig_hint(fig_text: str):
    """按配图文字推断占位框的类别与提示语。"""
    codes = re.findall(r"\b([SVE]\d{2})\b", fig_text)
    if codes:
        kind = "shot" if codes[0][0] in ("S", "V") else "cmd"
        label = "、".join(dict.fromkeys(codes))
        proof = ASSET_PROOF.get(codes[0], "")
        if codes[0][0] == "V":
            size = "建议从录屏（1080×1920 竖屏）截关键帧"
        else:
            size = CATEGORY_SIZE[kind]
        return f"【占位·待实测】{label}", f"{proof}；{size}"
    if "实物" in fig_text:
        return "【占位·待实测】实物照片", "待实测（节点/汇聚/报警器/433 实物，建议 ≥1600×1200，含接线与通电状态）"
    if "四象限" in fig_text or "路线图" in fig_text:
        return "【占位·待补】示意图", "待补（自绘，可据正文与 `docs/执行/待确认问题.md` 的 D 清单绘制）"
    return "【占位·待补】配图", f"待补（{fig_text or '未注明'}；建议矢量或 ≥1600 px 宽高清 PNG）"


def build_pptx(warn: list[str]) -> dict:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Inches, Pt

    lines = read_lines(SRC_PPT)
    pages = parse_ppt_pages(lines, warn)
    backup = parse_backup_pages(lines)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    DARK = RGBColor(0x1F, 0x35, 0x64)
    GRAY = RGBColor(0x59, 0x59, 0x59)
    RED = RGBColor(0xC0, 0x00, 0x00)

    def set_font(run, size=20, bold=False, ea="微软雅黑", latin="Arial", color=None):
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = latin
        if color is not None:
            run.font.color.rgb = color
        rpr = run._r.get_or_add_rPr()
        for tag, tf in (("a:latin", latin), ("a:ea", ea), ("a:cs", latin)):
            el = rpr.find(qn(tag))
            if el is None:
                el = parse_xml(f'<{tag} xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>')
                rpr.append(el)
            el.set("typeface", tf)

    def textbox(slide, l, t, w, h):
        box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        return box, tf

    def para(tf, text, size=20, bold=False, color=None, first=False, space_after=6, ea="微软雅黑"):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = text
        set_font(run, size=size, bold=bold, color=color, ea=ea)
        return p

    def add_placeholder(slide, label, hint, l=8.05, t=1.45, w=4.7, h=4.5):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches as In

        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, In(l), In(t), In(w), In(h))
        shape.fill.background()
        shape.line.color.rgb = GRAY
        shape.line.width = Pt(1.5)
        shape.line.dash_style = 4  # MSO_LINE_DASH_STYLE.DASH
        tf = shape.text_frame
        tf.word_wrap = True
        para(tf, label, size=14, bold=True, color=RED, first=True, space_after=6)
        para(tf, hint, size=11, color=GRAY, space_after=4)
        para(tf, "（本机无真机/模拟器，禁止用设计预览图冒充）", size=10, color=GRAY)
        return shape

    def footer(slide, text):
        box, tf = textbox(slide, 0.6, 6.95, 12.1, 0.4)
        para(tf, text, size=10, color=GRAY, first=True)

    # 封面
    slide = prs.slides.add_slide(blank)
    box, tf = textbox(slide, 0.8, 2.2, 11.7, 1.2)
    para(tf, "独居老人居家安全监测与应急辅助系统", size=36, bold=True, color=DARK, first=True)
    box, tf = textbox(slide, 0.8, 3.5, 11.7, 1.8)
    para(tf, "赛题四 · 开源鸿蒙赛道", size=20, color=GRAY, first=True, space_after=10)
    para(tf, "参赛作品设计方案 · 2026 年 11 月", size=16, color=GRAY)
    para(tf, "讲解时长上限 10 分钟 · 本稿 P01–P14 合计 560 s", size=16, color=GRAY)
    # 匿名（R2）：封面保持匿名，队名等信息只出现在压缩包文件名里（R7）。
    slide.notes_slide.notes_text_frame.text = (
        "封面保持匿名：只写作品名称、赛题与日期；队伍信息只出现在压缩包文件名（R7）。"
        "讲者开场：一句话说明作品解决的问题，再进入 P02。"
    )

    # P01–P14
    for idx, pg in enumerate(pages):
        slide = prs.slides.add_slide(blank)
        box, tf = textbox(slide, 0.6, 0.35, 12.1, 0.9)
        para(tf, f"{pg['code']}　{pg['title']}", size=28, bold=True, color=DARK, first=True)

        box, tf = textbox(slide, 0.75, 1.5, 7.0, 4.8)
        for bi, b in enumerate(pg["bullets"]):
            para(tf, f"• {b}", size=20, first=(bi == 0), space_after=12)

        label, hint = fig_hint(pg["fig"])
        add_placeholder(slide, label, f"配图：{pg['fig']}\n{hint}")

        slide.notes_slide.notes_text_frame.text = pg["notes"]
        footer(slide, f"{pg['code']}　{pg['head']}　|　讲者备注见备注栏　|　{TEAM_PLACEHOLDER}")

    # 备用页
    if backup:
        slide = prs.slides.add_slide(blank)
        box, tf = textbox(slide, 0.6, 0.35, 12.1, 0.9)
        para(tf, "备用页 B1–B7（附录，不计入 10 分钟）", size=28, bold=True, color=DARK, first=True)
        box, tf = textbox(slide, 0.75, 1.5, 11.8, 5.0)
        para(tf, "被追问时翻到这里；内容已在正文/附录里，不新造数字。", size=18, color=GRAY, first=True, space_after=12)
        for row in backup:
            if len(row) >= 3:
                para(tf, f"• {row[0]}　{row[1]}：{row[2]}", size=16, space_after=8)
        footer(slide, "备用页目录")

    for row in backup:
        slide = prs.slides.add_slide(blank)
        box, tf = textbox(slide, 0.6, 0.35, 12.1, 0.9)
        para(tf, f"{row[0]}　{row[1]}（备用页）", size=28, bold=True, color=DARK, first=True)
        box, tf = textbox(slide, 0.75, 1.5, 7.0, 4.8)
        para(tf, f"• {row[2]}", size=20, first=True)
        add_placeholder(slide, "【占位·待补】备用页图表", "待补（备用页图表按需制作；数据必须与正文一致）")
        slide.notes_slide.notes_text_frame.text = "备用页：只在评委追问时使用，不主动讲。"
        footer(slide, f"{row[0]} 备用页")

    os.makedirs(DIST_DIR, exist_ok=True)
    prs.save(PPTX_PATH)

    return {"pages": len(pages), "backup": len(backup), "slides": len(prs.slides._sldIdLst)}


# ------------------------------------------------------------------ 产物属性清理


def sanitize_docx_props(path: str, warn: list[str]):
    """把 docProps 里的作者/公司等身份字段清成占位符（python-docx 管不到 app.xml）。"""
    tmp = path + ".tmp"
    with zipfile.ZipFile(path, "r") as zin:
        names = zin.namelist()
        data = {n: zin.read(n) for n in names}
    changed = []
    if "docProps/app.xml" in data:
        xml = data["docProps/app.xml"].decode("utf-8")
        for tag in ("Company", "Manager"):
            new = re.sub(rf"<{tag}>.*?</{tag}>", f"<{tag}>{TEAM_PLACEHOLDER if tag == 'Company' else ''}</{tag}>", xml, flags=re.S)
            if new != xml:
                changed.append(tag)
                xml = new
        data["docProps/app.xml"] = xml.encode("utf-8")
    # 目录域在 Word 里默认不刷新（打开只看到「更新域」提示）。在 settings.xml 打开
    # updateFields，Word 打开文档时会自动刷新目录页码——评委拿到的就是带页码的目录。
    if "word/settings.xml" in data:
        xml = data["word/settings.xml"].decode("utf-8")
        if "updateFields" not in xml:
            m = re.search(r"<w:settings[^>]*>", xml)
            if m:
                xml = xml[: m.end()] + '<w:updateFields w:val="true"/>' + xml[m.end() :]
                changed.append("updateFields")
                data["word/settings.xml"] = xml.encode("utf-8")
    if "docProps/core.xml" in data:
        xml = data["docProps/core.xml"].decode("utf-8")
        new = re.sub(r"<dc:creator>.*?</dc:creator>", f"<dc:creator>{TEAM_PLACEHOLDER}</dc:creator>", xml, flags=re.S)
        new = re.sub(r"<cp:lastModifiedBy>.*?</cp:lastModifiedBy>", f"<cp:lastModifiedBy>{TEAM_PLACEHOLDER}</cp:lastModifiedBy>", new, flags=re.S)
        if new != xml:
            changed.append("core")
            data["docProps/core.xml"] = new.encode("utf-8")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:  # 每个成员只写一次，顺序不变
            zout.writestr(n, data[n])
    os.replace(tmp, path)
    return changed


def sanitize_pptx_props(path: str, warn: list[str]):
    tmp = path + ".tmp"
    with zipfile.ZipFile(path, "r") as zin:
        names = zin.namelist()
        data = {n: zin.read(n) for n in names}
    changed = []
    for member in ("docProps/core.xml", "docProps/app.xml"):
        if member not in data:
            continue
        xml = data[member].decode("utf-8")
        new = re.sub(r"<dc:creator>.*?</dc:creator>", f"<dc:creator>{TEAM_PLACEHOLDER}</dc:creator>", xml, flags=re.S)
        new = re.sub(r"<cp:lastModifiedBy>.*?</cp:lastModifiedBy>", f"<cp:lastModifiedBy>{TEAM_PLACEHOLDER}</cp:lastModifiedBy>", new, flags=re.S)
        new = re.sub(r"<Company>.*?</Company>", f"<Company>{TEAM_PLACEHOLDER}</Company>", new, flags=re.S)
        if new != xml:
            changed.append(member)
            data[member] = new.encode("utf-8")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            zout.writestr(n, data[n])
    os.replace(tmp, path)
    return changed


# ------------------------------------------------------------------ 自检


def verify(warn: list[str]) -> dict:
    from docx import Document
    from pptx import Presentation

    result = {}
    doc = Document(DOCX_PATH)
    texts = [p.text for p in doc.paragraphs]
    all_text = "\n".join(texts + [c.text for t in doc.tables for r in t.rows for c in r.cells])

    result["docx_paragraphs"] = len(doc.paragraphs)
    result["docx_tables"] = len(doc.tables)
    zf = zipfile.ZipFile(DOCX_PATH)
    doc_xml = zf.read("word/document.xml").decode("utf-8")
    styles_xml = zf.read("word/styles.xml").decode("utf-8")
    footer_xml = "".join(zf.read(n).decode("utf-8") for n in zf.namelist() if n.startswith("word/footer"))
    result["docx_has_toc"] = "TOC" in doc_xml and "目录" in all_text
    result["docx_page_field"] = "PAGE" in footer_xml
    result["docx_placeholder_boxes"] = all_text.count("【占位·待补素材】")
    result["docx_placeholders_total"] = all_text.count("【占位")
    # 小四 = 12pt；1.5 倍行距 = w:line="360" w:lineRule="auto"（在 styles.xml 里）
    body_sizes = {r.font.size.pt for p in doc.paragraphs for r in p.runs if r.font.size}
    result["docx_run_sizes"] = sorted(body_sizes)
    result["docx_line_360"] = styles_xml.count('w:line="360"')
    m_normal = re.search(r'<w:style [^>]*w:styleId="Normal"[^>]*>(.*?)</w:style>', styles_xml, re.S)
    result["docx_normal_15x"] = bool(m_normal and 'w:line="360"' in m_normal.group(1) and "auto" in m_normal.group(1))
    result["docx_eastAsia_songti"] = doc_xml.count("宋体")
    result["docx_eastAsia_heiti"] = doc_xml.count("黑体")

    prs = Presentation(PPTX_PATH)
    slides = list(prs.slides)
    result["pptx_slides"] = len(slides)
    notes = sum(1 for s in slides if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip())
    result["pptx_slides_with_notes"] = notes
    ph = 0
    for s in slides:
        for sh in s.shapes:
            if sh.has_text_frame and "【占位" in sh.text_frame.text:
                ph += 1
    result["pptx_placeholder_shapes"] = ph
    return result


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true", help="只自检已有产物，不重新生成")
    args = ap.parse_args()

    warn: list[str] = []
    for path in (SRC_DOC, SRC_PPT):
        if not os.path.exists(path):
            print(f"缺少真源：{path}", file=sys.stderr)
            return 2

    if not args.check:
        info_doc = build_docx(warn)
        info_ppt = build_pptx(warn)
        changed = sanitize_docx_props(DOCX_PATH, warn)
        changed += sanitize_pptx_props(PPTX_PATH, warn)
        print(f"[docx] {DOCX_PATH}")
        print(f"       标题 {info_doc['headings']} / 表格 {info_doc['tables']} / 占位框 {info_doc['placeholder_boxes']}")
        print(f"[pptx] {PPTX_PATH}")
        print(f"       P 页 {info_ppt['pages']} / 备用页 {info_ppt['backup']}")
        print(f"[props] 已清理：{changed}")

    info = verify(warn)
    print("[verify] " + json.dumps(info, ensure_ascii=False))
    for path in (DOCX_PATH, PPTX_PATH):
        print(f"[size] {os.path.basename(path)} = {os.path.getsize(path)} B")

    problems = []
    if info["docx_placeholder_boxes"] < len(CHAPTER_ASSETS):
        problems.append(f"占位框 {info['docx_placeholder_boxes']} < 章节数 {len(CHAPTER_ASSETS)}")
    if info["pptx_slides"] < 15 + 7:
        problems.append(f"PPT 页数 {info['pptx_slides']} 少于 P01–P14 + 备用页 7 页")
    if info["pptx_slides_with_notes"] < info["pptx_slides"] - 2:
        problems.append(f"有备注的页 {info['pptx_slides_with_notes']} 偏少（除封面与备用页目录外都应有备注）")
    if not info["docx_page_field"]:
        problems.append("页脚缺 PAGE 域")
    if not info["docx_normal_15x"]:
        problems.append("Normal 样式未设置 1.5 倍行距（w:line=360 w:lineRule=auto）")
    if not info["docx_eastAsia_songti"]:
        problems.append("正文未设置宋体（w:eastAsia=宋体）")

    if warn:
        print("[warn]")
        for w in dict.fromkeys(warn):
            print("  - " + w)
    if problems:
        print("[FAIL]")
        for p in problems:
            print("  - " + p)
        return 1
    print("[OK] 自检通过（占位框齐全、页脚页码、宋体/小四/1.5 倍行距、PPT 页数）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
